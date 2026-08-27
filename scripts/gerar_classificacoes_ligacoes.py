"""
PASSO 1 — Congela a taxonomia dos prompts de ligacao e gera a planilha
`ligacoes_genai/Classificacoes_ligacoes.xlsx` (uma aba por prompt).

Fonte de verdade da taxonomia (NAO o apresentacao/DICIONARIO_TECNICO.md, que
esta dessincronizado — ver nota no README):
  - scripts/analisar_gemini_julho.py  (pipeline de producao)
  - prompts/prompts_nivel1_por_ligacao.md
Os dois concordam entre si. O dicionario tecnico ainda tem a taxonomia ANTIGA
(B2=Storytelling, P5=Sinistro, FC5=Inaplicavel, AB so ate AB4, desfecho="fechou").

Estrutura de cada aba:
  grupo | codigo | rotulo | descricao | peso
- `grupo`: agrupa os campos de um mesmo prompt (ex: em P1: "desfecho",
  "intencao_entrada", "comportamento").
- `peso`: so preenchido para os comportamentos B1-B10 de P1 (entra no score).

Reproduzivel e versionavel — rode de novo sempre que a taxonomia mudar.
Execute:  python scripts/gerar_classificacoes_ligacoes.py
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SAIDA = Path(__file__).parent.parent / "ligacoes_genai" / "Classificacoes_ligacoes.xlsx"

# ---------------------------------------------------------------------------
# Taxonomia — cada entrada: (grupo, codigo, rotulo, descricao, peso)
# ---------------------------------------------------------------------------

P1 = [
    # --- Camada 1: desfecho da ligacao ---
    ("desfecho", "fechou_novo", "Fechou nova locacao",
     "Nova locacao/contrato que nao existia antes.", None),
    ("desfecho", "fechou_renovacao", "Fechou renovacao",
     "Cliente renovou ou prorrogou contrato existente.", None),
    ("desfecho", "fechou_upsell", "Fechou upsell",
     "Cliente adicionou produto/servico que nao tinha.", None),
    ("desfecho", "interessou_nao_fechou", "Interessou, nao fechou",
     "Interesse demonstrado, mas nao concluiu na ligacao.", None),
    ("desfecho", "nao_era_venda", "Nao era venda",
     "Ligacao 100% operacional/suporte/administrativa.", None),
    ("desfecho", "sem_conteudo", "Sem conteudo",
     "Caixa postal, ligacao muda, conexao perdida, sem conversa identificavel.", None),
    # --- Camada 2: intencao de entrada ---
    ("intencao_entrada", "cliente_queria", "Cliente queria",
     "Cliente ligou com intencao clara de alugar/renovar.", None),
    ("intencao_entrada", "vendedor_puxou", "Vendedor puxou",
     "Vendedor ligou proativamente e abriu o assunto comercial.", None),
    ("intencao_entrada", "misto", "Misto",
     "Cliente ligou por outro motivo e o vendedor aproveitou.", None),
    ("intencao_entrada", "indefinido", "Indefinido",
     "Impossivel determinar a intencao de entrada.", None),
    # --- Camada 3: comportamentos do vendedor (SIM/NAO/NA) ---
    ("comportamento", "B1", "Diagnostico antes do produto",
     "Fez pergunta de diagnostico sobre a situacao do cliente antes de apresentar o produto.", 2),
    ("comportamento", "B2", "Ancora de preco alto",
     "Apresentou o produto mais caro primeiro, como ancora de preco.", 1),
    ("comportamento", "B3", "Usou historico do cliente",
     "Trouxe um dado concreto do historico/contrato do cliente para a conversa.", 2),
    ("comportamento", "B4", "Avisou risco ao cliente",
     "Avisou o cliente de algo que poderia prejudica-lo (custo, prazo, regra).", 2),
    ("comportamento", "B5", "Upsell conectado a dor",
     "Tentou upsell conectado a uma dor que o cliente revelou nesta ligacao "
     "(nao pitch generico).", 2),
    ("comportamento", "B6", "Urgencia com fato verificavel",
     "Criou urgencia com fato verificavel (prazo real, mudanca de preco, disponibilidade).", 3),
    ("comportamento", "B7", "Vai batalhar internamente",
     "Disse que vai batalhar/negociar internamente pelo cliente.", 1),
    ("comportamento", "B8", "Ligou proativamente",
     "Ligou proativamente por motivo comercial ou de retencao.", 1),
    ("comportamento", "B9", "Conhecimento pessoal do cliente",
     "Demonstrou conhecimento pessoal do cliente alem do contrato.", 1),
    ("comportamento", "B10", "Sondou concorrente/frota propria",
     "Perguntou se o cliente usa outra locadora ou tem frota propria.", 3),
]
P1_NOTA = ("score_comercial = soma dos pesos dos comportamentos marcados SIM. "
           "Maximo teorico = 18. tentativa_comercial: SIM/NAO/desconhecido.")

P2 = [
    ("tipo", "nova_venda", "Nova venda", "Primeiro contrato, cotacao de nova frota.", None),
    ("tipo", "renovacao", "Renovacao", "Prorrogacao, novo periodo de contrato existente.", None),
    ("tipo", "upsell", "Upsell", "Telemetria, protecao, categoria superior, km adicional.", None),
    ("tipo", "retencao", "Retencao", "Queda de volume, risco de churn, cliente menciona concorrente.", None),
    ("tipo", "suporte_operacional", "Suporte operacional", "Portal, acesso, sistema, reserva com problema.", None),
    ("tipo", "pos_venda_sinistro", "Pos-venda / sinistro", "Acidente, avaria, oficina, franquia de dano.", None),
    ("tipo", "pos_venda_manutencao", "Pos-venda / manutencao", "Carro em manutencao, substituto, prazo de reparo.", None),
    ("tipo", "cobranca", "Cobranca", "Faturas atrasadas, boleto, bloqueio de frota.", None),
    ("tipo", "onboarding", "Onboarding", "Cadastro de condutor, biometria, troca de condutor.", None),
    ("tipo", "duvida_contrato", "Duvida de contrato", "Multa, km excedente, regras, diferenca mensal vs anual.", None),
    ("tipo", "misto", "Misto", "Dois ou mais tipos na mesma ligacao (subtipo = ambos separados por '+').", None),
    ("tipo", "sem_conteudo", "Sem conteudo", "Caixa postal, ligacao muda, conexao perdida.", None),
    ("quem_iniciou", "vendedor", "Vendedor iniciou", "O vendedor fez a ligacao.", None),
    ("quem_iniciou", "cliente", "Cliente iniciou", "O cliente fez a ligacao.", None),
    ("quem_iniciou", "indefinido", "Indefinido", "Nao da para determinar quem iniciou.", None),
]
P2_NOTA = ("tentativa_comercial = SIM se o vendedor tentou algo comercial, mesmo em "
           "ligacao de suporte. Se mistura suporte + tentativa comercial: tipo = misto.")

P3 = [
    ("problema", "P1", "Acesso/sistema", "Portal fora do ar, senha, biometria, dois portais distintos.", None),
    ("problema", "P2", "Faturamento/cobranca", "Boleto errado, fatura para terceiro, valor incorreto.", None),
    ("problema", "P3", "Multa/infracao", "Multa no condutor errado, prazo de indicacao, contestacao.", None),
    ("problema", "P4", "Manutencao/substituto", "Carro em oficina, demora, sem substituto disponivel.", None),
    ("problema", "P5", "Condutor/cadastro", "Troca de condutor, demitido com acesso, habilitacao.", None),
    ("problema", "P6", "Km excedente", "Cobranca incorreta, plano errado, negociacao de desconto.", None),
    ("problema", "P7", "Disponibilidade", "Carro reservado indisponivel, categoria esgotada.", None),
    ("problema", "P8", "Franquia/agencia", "Processo diferente em franqueada, fatura nao baixada.", None),
    ("problema", "P9", "Sinistro", "Acidente, franquia de dano, mudanca de oficina.", None),
    ("problema", "P10", "Reserva/sistema",
     "Reserva nao aparece no sistema, status incorreto, erro ao criar/editar reserva, "
     "bug reproduzivel do fluxo de reserva. NAO marcar quando for so duvida de navegacao "
     "sem erro real, nem quando 'reserva' so aparece de passagem e o assunto real e outro.", None),
    ("problema", "P11", "Outro", "Problema real que nao se encaixa em P1-P10 (descrever na observacao).", None),
    ("problema", "P12", "Manipulacao de pesquisa/NPS",
     "Vendedor pede/dita a nota da pesquisa de satisfacao. Problema de integridade do "
     "indicador, nao do cliente. Marcar quem_relatou=vendedor_detectou, "
     "foi_resolvido_na_ligacao=NAO.", None),
    ("quem_relatou", "cliente", "Cliente relatou", "O cliente trouxe o problema.", None),
    ("quem_relatou", "vendedor_detectou", "Vendedor detectou", "O vendedor identificou o problema.", None),
    ("quem_relatou", "nenhum", "Nenhum", "Ligacao sem problema (problemas_identificados = nenhum).", None),
    ("foi_resolvido_na_ligacao", "SIM", "Resolvido", "Problema resolvido dentro da propria ligacao.", None),
    ("foi_resolvido_na_ligacao", "NAO", "Nao resolvido", "Ficou em aberto ao final da ligacao.", None),
    ("foi_resolvido_na_ligacao", "PARCIAL", "Parcial", "Parte resolvida, parte em aberto.", None),
    ("escalou_para", "0800", "0800", "Encaminhado para a central 0800.", None),
    ("escalou_para", "supervisora", "Supervisora", "Encaminhado para a supervisora/lideranca.", None),
    ("escalou_para", "IT", "IT", "Encaminhado para o time de tecnologia.", None),
    ("escalou_para", "nao_escalou", "Nao escalou", "Vendedor tratou sem escalar.", None),
]
P3_NOTA = ("problemas_identificados: codigos separados por '+' (ex: P1+P3), ou 'nenhum'. "
           "So marcar um codigo se o problema for tema central/relevante da ligacao.")

P5 = [
    ("challenger", "CH1", "Custo oculto revelado",
     "Mostrou um custo que o cliente nao havia calculado (ex: diaria x dias > mensal).", None),
    ("challenger", "CH2", "ROI calculado",
     "Calculou o retorno financeiro do produto para aquele cliente especifico.", None),
    ("challenger", "CH3", "Urgencia por regra real",
     "Criou pressao com dado real: prazo que expira, preco que sobe, disponibilidade limitada.", None),
    ("challenger", "CH4", "Concorrente superado com dado",
     "Usou dado comparativo concreto para se diferenciar da concorrencia "
     "(ex: taxa de operacao, cobertura).", None),
    ("challenger", "CH5", "Dor conectada ao produto",
     "Ligou uma dor declarada pelo cliente diretamente a uma funcionalidade do produto.", None),
    ("challenger", "CH6", "Antecipacao de problema",
     "Alertou sobre um problema futuro que o cliente ainda nao havia percebido "
     "(ex: contrato vencendo).", None),
    ("challenger", "CH7", "Necessidade nao declarada",
     "Descobriu, perguntando, uma necessidade que o cliente nao havia declarado "
     "(ex: frota propria, caminhoes).", None),
    ("teve_challenger", "SIM", "Teve challenger", "Houve ao menos um momento Challenger na ligacao.", None),
    ("teve_challenger", "NAO", "Nao teve", "Nenhum comportamento Challenger. Demais campos = nenhum.", None),
    ("resultado_imediato", "cliente_aceitou", "Cliente aceitou", "Reacao positiva imediata do cliente.", None),
    ("resultado_imediato", "cliente_refletiu", "Cliente refletiu", "Cliente ponderou, sem aceitar na hora.", None),
    ("resultado_imediato", "sem_reacao_visivel", "Sem reacao visivel", "Nenhuma reacao perceptivel do cliente.", None),
    ("qualidade", "alta", "Alta",
     "Dado especifico e verificavel, conectado a algo que o cliente disse, com reacao visivel.", None),
    ("qualidade", "media", "Media", "Challenger presente, mas generico ou pouco conectado.", None),
    ("qualidade", "baixa", "Baixa", "Tentativa fraca, sem dado concreto ou sem conexao com a fala do cliente.", None),
]
P5_NOTA = ("codigos_challenger: um ou mais CH separados por ';' (ex: CH1;CH6), ou 'nenhum'. "
           "Criterio de qualidade alta: ver treinamento/GUIA_CHALLENGER_BEM_FEITO.md.")

P9 = [
    ("abertura", "AB1", "Contextualizada",
     "Vendedor diz por que esta ligando / referencia o historico antes de pedir atencao.", None),
    ("abertura", "AB2", "Relacional",
     "Comeca com vinculo pessoal, pergunta sobre o cliente antes do negocio.", None),
    ("abertura", "AB3", "Generica", "'Em que posso ajudar?' / 'tudo bem?' sem contexto nem personalizacao.", None),
    ("abertura", "AB4", "Reativa", "So responde ao que o cliente traz — cliente tomou a iniciativa.", None),
    ("abertura", "AB5", "De protecao", "Abre antecipando um problema que o cliente ainda nao viu.", None),
    ("fechamento", "FC1", "Compromisso duplo com prazo",
     "Define o que o vendedor E o cliente vao fazer, com prazo/data. O fechamento mais forte.", None),
    ("fechamento", "FC2", "Proximo passo so do vendedor",
     "Define so a acao do vendedor ('vou te mandar a proposta'), sem compromisso/prazo do cliente.", None),
    ("fechamento", "FC3", "Convite generico",
     "'Qualquer coisa me chama' / 'fico a disposicao' — sem data, sem proximo passo definido.", None),
    ("fechamento", "FC4", "Aberto", "Encerra sem definir nada.", None),
    ("fechamento", "FC5", "Diretivo",
     "Da instrucao clara ao cliente ('nao paga essa fatura ainda', 'acessa o portal e faz X').", None),
    ("usou_nome", "SIM", "Usou o nome", "Vendedor chamou o cliente pelo nome.", None),
    ("usou_nome", "NAO", "Nao usou", "Vendedor nao usou o nome do cliente.", None),
    ("proximo_passo_claro", "SIM", "Proximo passo claro", "Ficou definido o que acontece depois.", None),
    ("proximo_passo_claro", "NAO", "Sem proximo passo", "Nao ficou definido nada para depois.", None),
    ("prazo_definido", "SIM", "Prazo definido", "Ha data/prazo concreto.", None),
    ("prazo_definido", "NAO", "Sem prazo", "Nenhum prazo concreto.", None),
]
P9_NOTA = ("desfecho_p9: reaproveita o desfecho da ligacao (mesma taxonomia de P1). "
           "Analise a abertura pelas 3-5 primeiras falas do vendedor e o fechamento pelas 3-5 ultimas.")

ABAS = [
    ("P1", "Desfecho e Comportamentos Comerciais", P1, P1_NOTA),
    ("P2", "Tipo de Ligacao", P2, P2_NOTA),
    ("P3", "Problemas Recorrentes", P3, P3_NOTA),
    ("P5", "Comportamentos Challenger", P5, P5_NOTA),
    ("P9", "Abertura e Fechamento", P9, P9_NOTA),
]

# ---------------------------------------------------------------------------
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(color="FFFFFF", bold=True)
NOTA_FONT = Font(italic=True, color="555555")
GRP_FONT = Font(bold=True)
COLS = ["grupo", "codigo", "rotulo", "descricao", "peso"]
LARGURAS = [22, 12, 34, 90, 6]


def build():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for cod, nome, linhas, nota in ABAS:
        ws = wb.create_sheet(cod)
        ws["A1"] = f"{cod} — {nome}"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = nota
        ws["A2"].font = NOTA_FONT
        ws.append([])

        hdr_row = ws.max_row + 1
        ws.append(COLS)
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(row=hdr_row, column=c)
            cell.fill = HEAD_FILL
            cell.font = HEAD_FONT

        grupo_ant = None
        for grupo, codigo, rotulo, descricao, peso in linhas:
            ws.append([grupo, codigo, rotulo, descricao, peso if peso is not None else ""])
            r = ws.max_row
            if grupo != grupo_ant:
                ws.cell(row=r, column=1).font = GRP_FONT
                grupo_ant = grupo
            ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")

        for i, w in enumerate(LARGURAS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    print(f"OK -> {SAIDA}")
    for cod, nome, linhas, _ in ABAS:
        grupos = {}
        for g, c, *_ in linhas:
            grupos.setdefault(g, 0)
            grupos[g] += 1
        resumo = " | ".join(f"{g}:{n}" for g, n in grupos.items())
        print(f"  {cod}: {len(linhas)} linhas  ({resumo})")


if __name__ == "__main__":
    build()
