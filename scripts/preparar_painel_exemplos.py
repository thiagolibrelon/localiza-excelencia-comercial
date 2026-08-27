"""
Prepara o painel de 5 agentes independentes sobre Exemplos_ligacoes_CANDIDATOS.xlsx.

Gera:
  ligacoes_genai/_panel/candidatos.json  -> o que os agentes recebem (SEM o rotulo
                                            do Gemini — classificacao as cegas)
  ligacoes_genai/_panel/_gabarito.json   -> rotulo do Gemile, so pra consolidacao

Cada candidato ganha um id estavel: {prompt}-{NN} (ordem das linhas do xlsx).

Execute:  python scripts/preparar_painel_exemplos.py
"""
from pathlib import Path
import json
import openpyxl

RAIZ = Path(__file__).parent.parent
FONTE = RAIZ / "ligacoes_genai" / "Exemplos_ligacoes_CANDIDATOS.xlsx"
PANEL = RAIZ / "ligacoes_genai" / "_panel"


def build():
    wb = openpyxl.load_workbook(FONTE)
    candidatos, gabarito = [], {}

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        seq = 0
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[8]:
                continue
            seq += 1
            cid = f"{sheet}-{seq:02d}"
            candidatos.append({
                "id": cid,
                "prompt": sheet,
                "alvo_amostragem": row[2],     # so contexto (por que a linha foi puxada)
                "cliente_id": row[3],
                "data": row[4],
                "trecho_transcricao": row[8],
            })
            gabarito[cid] = row[7]             # classificacao_gemini

    PANEL.mkdir(parents=True, exist_ok=True)
    (PANEL / "candidatos.json").write_text(
        json.dumps(candidatos, ensure_ascii=False, indent=1), encoding="utf-8")
    (PANEL / "_gabarito.json").write_text(
        json.dumps(gabarito, ensure_ascii=False, indent=1), encoding="utf-8")

    por_prompt = {}
    for c in candidatos:
        por_prompt[c["prompt"]] = por_prompt.get(c["prompt"], 0) + 1
    print(f"OK -> {PANEL/'candidatos.json'}  ({len(candidatos)} candidatos)")
    for p, n in sorted(por_prompt.items()):
        print(f"  {p}: {n}")


if __name__ == "__main__":
    build()
